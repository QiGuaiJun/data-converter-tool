from __future__ import annotations

import atexit
import datetime as _dt
import json
import os
import shutil
import tempfile
from pathlib import Path
from io import BytesIO

import pytest
from openpyxl import load_workbook

# 隔离测试环境：数据/上传/导出落临时目录，不触碰真实 data/ uploads/ exports/
_tmp = tempfile.mkdtemp(prefix="dc_test_")
os.environ["DATA_DIR"] = str(Path(_tmp) / "data")
os.environ["UPLOADS_DIR"] = str(Path(_tmp) / "uploads")
os.environ["EXPORTS_DIR"] = str(Path(_tmp) / "exports")

import server

# sqlite3 连接作为 with 语句使用时只提交事务、不 close，句柄会一直占着临时库文件；
# Windows 上这会挡住隔离目录的删除，把 %TEMP%/dc_test_* 越攒越多。这里记录句柄，
# 进程退出时统一关闭再删目录，保证跑完 %TEMP% 零残留。
_TRACKED_CONNECTIONS: list[object] = []
_ORIGINAL_CONNECT_DB = server.connect_db


def _tracking_connect_db():
    conn = _ORIGINAL_CONNECT_DB()
    _TRACKED_CONNECTIONS.append(conn)
    return conn


server.connect_db = _tracking_connect_db


@atexit.register
def _remove_isolated_tmp() -> None:
    for conn in _TRACKED_CONNECTIONS:
        try:
            conn.close()
        except Exception:  # pragma: no cover - 关闭失败不影响测试结论
            pass
    shutil.rmtree(_tmp, ignore_errors=True)


# MySQL 用例的连接参数：默认本机，可用 DC_TEST_MYSQL_* 环境变量覆盖。
# 与 test_p2_fixes.py 保持同一套约定，避免两个文件跑在不同库上。
MYSQL_FIELDS = {
    "targetDbType": "mysql",
    "dbHost": os.environ.get("DC_TEST_MYSQL_HOST", "127.0.0.1"),
    "dbPort": os.environ.get("DC_TEST_MYSQL_PORT", "3306"),
    "dbUser": os.environ.get("DC_TEST_MYSQL_USER", "root"),
    "dbPassword": os.environ.get("DC_TEST_MYSQL_PASSWORD", "123456"),
    "dbName": os.environ.get("DC_TEST_MYSQL_DB", "dc_p2_test"),
}

# 专用测试库 -- 表名统一加 fixt_ 前缀，只在本文件用例内创建和销毁。
FIXT_TABLE = "fixt_export_people"
FIXT_IMPORT_TABLE = "fixt_export_people_conn"


def mysql_available() -> bool:
    """探测测试用 MySQL 是否可连接，并确保目标库存在。"""
    try:
        import pymysql

        conn = pymysql.connect(
            host=MYSQL_FIELDS["dbHost"],
            port=int(MYSQL_FIELDS["dbPort"]),
            user=MYSQL_FIELDS["dbUser"],
            password=MYSQL_FIELDS["dbPassword"],
            connect_timeout=3,
        )
        with conn.cursor() as cursor:
            cursor.execute(
                f"create database if not exists {MYSQL_FIELDS['dbName']} default character set utf8mb4"
            )
        conn.commit()
        conn.close()
        return True
    except Exception:
        return False


MYSQL_READY = mysql_available()


def sqlite_fixture() -> None:
    with server.connect_db() as conn:
        conn.execute("drop table if exists export_people")
        conn.execute("create table export_people (name text, amount integer, city text)")
        conn.executemany(
            "insert into export_people values (?, ?, ?)",
            [("Alice", 10, "北京"), ("Bob", 20, "上海"), ("Carol", 30, "北京")],
        )


def mysql_ensure_fixture_tables() -> None:
    """在测试库里建 fixt_ 前缀的表（若可用），供 MySQL 导出用例使用。"""
    conn = server.connect_target_db(dict(MYSQL_FIELDS))
    try:
        with conn.cursor() as cursor:
            cursor.execute(f"drop table if exists {FIXT_TABLE}")
            cursor.execute(
                f"create table {FIXT_TABLE} (name varchar(32), amount int, city varchar(32))"
            )
            cursor.executemany(
                f"insert into {FIXT_TABLE} values (%s, %s, %s)",
                [("Alice", 10, "北京"), ("Bob", 20, "上海"), ("Carol", 30, "北京")],
            )
        conn.commit()
    finally:
        conn.close()


def mysql_drop_fixture_tables() -> None:
    """清理本文件在测试库里创建的所有表，跑完不留下 fixt_ 残留。"""
    conn = server.connect_target_db(dict(MYSQL_FIELDS))
    try:
        with conn.cursor() as cursor:
            for table in (FIXT_TABLE, FIXT_IMPORT_TABLE):
                cursor.execute(f"drop table if exists {table}")
        conn.commit()
    finally:
        conn.close()


def assert_xlsx(path: str, expected_rows: int) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    assert len(rows) == expected_rows


def test_sqlite_exports() -> None:
    sqlite_fixture()
    result = server.run_export_job(
        {
            "targetDbType": "sqlite",
            "items": [{"type": "table", "table": "export_people", "name": "export_people"}],
            "extension": "xlsx",
            "outputName": "qa_export_people",
            "sheetName": "People",
            "exportMode": "workbook",
            "headerMode": "field",
            "exportFields": "name,amount",
            "whereClause": "amount >= 20",
            "rowHeight": "22",
            "columnWidth": "18",
            "fontName": "Arial",
            "fontSize": "11",
            "addBorder": "true",
            "lockHeader": "true",
        }
    )
    assert result["rows"] == 2
    assert_xlsx(result["files"][0], 3)

    csv_result = server.run_export_job(
        {
            "targetDbType": "sqlite",
            "items": [{"type": "query", "name": "query_amount", "sql": "select name, amount from export_people order by amount"}],
            "extension": "csv",
            "outputName": "qa_export_people_csv",
            "headerMode": "field",
            "encoding": "utf-8",
            "delimiter": ",",
        }
    )
    assert Path(csv_result["files"][0]).read_text(encoding="utf-8").startswith("name,amount")

    json_result = server.run_export_job(
        {
            "targetDbType": "sqlite",
            "items": [{"type": "table", "table": "export_people", "name": "export_people"}],
            "extension": "json",
            "outputName": "qa_export_people_json",
            "splitField": "city",
        }
    )
    assert len(json_result["files"]) == 2
    assert json.loads(Path(json_result["files"][0]).read_text(encoding="utf-8"))

    xml_result = server.run_export_job(
        {
            "targetDbType": "sqlite",
            "items": [{"type": "query", "name": "query_xml", "sql": "select name from export_people where city = '北京'"}],
            "extension": "xml",
            "outputName": "qa_export_people_xml",
        }
    )
    assert Path(xml_result["files"][0]).read_text(encoding="utf-8").startswith("<?xml")


def _ensure_mysql_connection_id() -> str | None:
    """在隔离的 sqlite 里，用产品自身落库路径写入一条可用的 MySQL 连接记录。

    走 ``server.normalize_connection_payload``（校验）+ ``server.encode_secret``（Fernet 加密）
    再 insert 的落库逻辑，与 ``server.handle_connection_save`` 一致，不绕过加密/校验直接插明文。
    返回 connectionId；仅当本机 MySQL 确实不可连时返回 None（调用方据此 pytest.skip）。
    """
    global _MYSQL_CONN_ID
    if _MYSQL_CONN_ID is not None:
        return _MYSQL_CONN_ID

    if not MYSQL_READY:
        return None

    payload = {
        "id": "conn_qa_m2058",
        "name": "QA M2-058 / M12-016",
        "dbType": "mysql",
        "host": MYSQL_FIELDS["dbHost"],
        "port": int(MYSQL_FIELDS["dbPort"]),
        "user": MYSQL_FIELDS["dbUser"],
        "password": MYSQL_FIELDS["dbPassword"],
        "database": MYSQL_FIELDS["dbName"],
        "charset": "utf8mb4",
    }
    record = server.normalize_connection_payload(payload)
    now = _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with server.connect_db() as conn:
        conn.execute(
            """
            insert into _db_connections (
                id, name, db_type, host, port, user_name, password, db_name, charset,
                ssl_enabled, ssl_ca, ssl_cert, ssl_key, created_at, updated_at
            ) values (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            on conflict(id) do update set
                name=excluded.name, db_type=excluded.db_type, host=excluded.host,
                port=excluded.port, user_name=excluded.user_name, password=excluded.password,
                db_name=excluded.db_name, charset=excluded.charset, updated_at=excluded.updated_at
            """,
            (
                record["id"], record["name"], record["db_type"], record["host"],
                record["port"], record["user_name"],
                server.encode_secret(str(record["password"])),
                record["db_name"], record["charset"], record["ssl_enabled"],
                record["ssl_ca"], record["ssl_cert"], record["ssl_key"], now, now,
            ),
        )
    # 真断言：保存后能读回并解密出原始密码，证明加密/解密落库链路真实可用。
    loaded = server.load_saved_connection(record["id"])
    assert loaded["dbPassword"] == MYSQL_FIELDS["dbPassword"], (
        "连接密码保存后无法解密回原始明文，加密/落库路径异常"
    )
    _MYSQL_CONN_ID = record["id"]
    return _MYSQL_CONN_ID


_MYSQL_CONN_ID: str | None = None


def test_mysql_export_if_available() -> None:
    """连接式（connectionId）导入 + 导出真实覆盖（M2-058 / M12-016）。

    先在隔离库里用产品自身落库路径创建一条指向 127.0.0.1:3306/dc_p2_test 的 MySQL 连接，
    再以 connectionId 完成「导入 CSV -> 目标 MySQL 表」与「从目标 MySQL 表导出 xlsx」两条链路，
    并对导入行数、导出行数、xlsx 列头与数据做实质断言。仅在本机 MySQL 确实连不上时 pytest.skip。
    """
    connection_id = _ensure_mysql_connection_id()
    if connection_id is None:
        pytest.skip("本机 MySQL 不可连，跳过连接式导入/导出用例")

    fields = {
        "connectionId": connection_id,
        "importMode": "rebuild",
        "tableName": FIXT_IMPORT_TABLE,
        "mapping": json.dumps(
            [
                {"sourceIndex": 0, "target": "Name", "enabled": True, "defaultValue": "", "matchKey": True},
                {"sourceIndex": 1, "target": "Amount", "enabled": True, "defaultValue": "0", "matchKey": False},
            ]
        ),
        "tableCase": "lower",
        "fieldCase": "lower",
        "commitMode": "once",
    }
    upload = Path(_tmp) / "uploads" / f"{FIXT_IMPORT_TABLE}.csv"
    upload.parent.mkdir(parents=True, exist_ok=True)
    upload.write_text("Name,Amount\nA,1\nB,2\n", encoding="utf-8-sig")
    import_result = server.import_uploaded_file(server.UploadedFile(f"{FIXT_IMPORT_TABLE}.csv", upload), fields)
    # 真断言：连接式导入确实写入了目标 MySQL 表。
    assert import_result["rowsWritten"] == 2, import_result

    try:
        result = server.run_export_job(
            {
                "connectionId": connection_id,
                "targetDbType": "mysql",
                "items": [{"type": "table", "table": FIXT_IMPORT_TABLE, "name": FIXT_IMPORT_TABLE}],
                "extension": "xlsx",
                "outputName": "qa_mysql_export_people",
                "sheetName": "MySQL",
            }
        )
        # 真断言：连接式导出返回正确行数 + 生成的 xlsx 存在且列头/数据正确。
        assert result["rows"] == 2, result
        workbook = load_workbook(result["files"][0])
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        assert len(rows) == 3, rows  # 表头 + 2 行数据
        assert [str(c).lower() for c in rows[0]] == ["name", "amount"], rows[0]
        assert list(rows[1]) == ["A", 1], rows[1]
    finally:
        conn = server.connect_target_db({"connectionId": connection_id})
        try:
            with conn.cursor() as cursor:
                cursor.execute(f"drop table if exists {FIXT_IMPORT_TABLE}")
            conn.commit()
        finally:
            conn.close()


@pytest.mark.skipif(not MYSQL_READY, reason="本机 MySQL 不可用")
def test_mysql_export_with_direct_fields() -> None:
    """真实 MySQL 导出覆盖：不依赖 _db_connections，直接用环境变量参数连库。

    fixt_export_people 只存在于 MySQL（SQLite 里没有同名表），因此一旦
    targetDbType 解析失败、静默回退到 SQLite，本用例必然报“no such table”，
    不存在“在 SQLite 上跑通然后算 MySQL 覆盖”的假象。
    """
    # 非 pytest 运行器（脚本模式）不会执行 skipif 标记，这里再挡一层。
    if not MYSQL_READY:
        pytest.skip("本机 MySQL 不可用")

    mysql_ensure_fixture_tables()
    try:
        # 先在 MySQL 侧确认数据确实落在测试库里
        sources = server.export_sources(dict(MYSQL_FIELDS))
        entry = next(item for item in sources if item["name"] == FIXT_TABLE)
        assert entry["rows"] == 3

        xlsx_result = server.run_export_job(
            {
                **MYSQL_FIELDS,
                "items": [{"type": "table", "table": FIXT_TABLE, "name": FIXT_TABLE}],
                "extension": "xlsx",
                "outputName": "qa_fixt_mysql_export",
                "sheetName": "MySQL",
                "headerMode": "field",
            }
        )
        assert xlsx_result["rows"] == 3
        assert_xlsx(xlsx_result["files"][0], 4)

        csv_result = server.run_export_job(
            {
                **MYSQL_FIELDS,
                "items": [
                    {
                        "type": "query",
                        "name": "fixt_mysql_query",
                        "sql": f"select name, amount from {FIXT_TABLE} order by amount",
                    }
                ],
                "extension": "csv",
                "outputName": "qa_fixt_mysql_export_csv",
                "headerMode": "field",
                "encoding": "utf-8",
                "delimiter": ",",
            }
        )
        assert csv_result["rows"] == 3
        csv_lines = Path(csv_result["files"][0]).read_text(encoding="utf-8").splitlines()
        assert csv_lines[0] == "name,amount"
        assert csv_lines[1] == "Alice,10"
        assert csv_lines[3] == "Carol,30"
    finally:
        mysql_drop_fixture_tables()


def test_large_streaming_export() -> None:
    with server.connect_db() as conn:
        conn.execute("drop table if exists export_large_people")
        conn.execute("create table export_large_people (id integer, name text, amount integer)")
        conn.executemany(
            "insert into export_large_people values (?, ?, ?)",
            ((index, f"name_{index}", index % 100) for index in range(20000)),
        )

    xlsx_result = server.run_export_job(
        {
            "targetDbType": "sqlite",
            "items": [{"type": "table", "table": "export_large_people", "name": "export_large_people"}],
            "extension": "xlsx",
            "outputName": "qa_large_streaming_xlsx",
            "sheetName": "Large",
            "exportMode": "workbook",
            "headerMode": "field",
        }
    )
    assert xlsx_result["rows"] == 20000
    assert_xlsx(xlsx_result["files"][0], 20001)

    csv_result = server.run_export_job(
        {
            "targetDbType": "sqlite",
            "items": [{"type": "table", "table": "export_large_people", "name": "export_large_people"}],
            "extension": "csv",
            "outputName": "qa_large_streaming_csv",
            "headerMode": "field",
        }
    )
    assert csv_result["rows"] == 20000
    assert Path(csv_result["files"][0]).read_text(encoding="utf-8").splitlines()[0] == "id,name,amount"


class FakeDownloadHandler:
    def __init__(self) -> None:
        self.headers: list[tuple[str, str]] = []
        self.status = None
        self.wfile = BytesIO()

    def send_response(self, status: int) -> None:
        self.status = status

    def send_header(self, key: str, value: str) -> None:
        value.encode("latin-1")
        self.headers.append((key, value))

    def end_headers(self) -> None:
        pass


def test_chinese_filename_download_header() -> None:
    path = server.EXPORTS / "\u4e2d\u6587\u5bfc\u51fa\u6587\u4ef6.xlsx"
    path.write_bytes(b"demo")
    handler = FakeDownloadHandler()
    server.ImportPrototypeHandler.handle_export_download(handler, "name=%E4%B8%AD%E6%96%87%E5%AF%BC%E5%87%BA%E6%96%87%E4%BB%B6.xlsx")
    assert handler.status == 200
    disposition = dict(handler.headers)["Content-Disposition"]
    assert disposition == "attachment; filename=\"export.xlsx\"; filename*=UTF-8''%E4%B8%AD%E6%96%87%E5%AF%BC%E5%87%BA%E6%96%87%E4%BB%B6.xlsx"



if __name__ == "__main__":
    checks = [
        ("test_sqlite_exports", test_sqlite_exports),
        ("test_mysql_export_if_available", test_mysql_export_if_available),
        ("test_mysql_export_with_direct_fields", test_mysql_export_with_direct_fields),
        ("test_large_streaming_export", test_large_streaming_export),
        ("test_chinese_filename_download_header", test_chinese_filename_download_header),
    ]
    passed: list[str] = []
    skipped: list[str] = []
    for name, check in checks:
        try:
            check()
        except pytest.skip.Exception as exc:  # pytest.skip 抛的是 BaseException 子类，必须显式捕获
            reason = str(exc) or "未提供跳过原因"
            print(f"[skip] {name}：{reason}")
            skipped.append(f"{name}：{reason}")
        else:
            passed.append(name)

    # 汇总必须如实反映 MySQL 覆盖是被执行了还是被跳过了，不能笼统宣称"通过"。
    mysql_checks = [name for name, _ in checks if "mysql" in name]
    mysql_passed = [name for name in passed if name in mysql_checks]
    mysql_skipped = [item for item in skipped if item.split("：", 1)[0] in mysql_checks]
    print(f"export engine checks passed（{len(passed)} 项通过，{len(skipped)} 项跳过）")
    if mysql_passed:
        print(f"MySQL 导出路径已覆盖并执行通过：{', '.join(mysql_passed)}")
    if mysql_skipped:
        print(f"MySQL 导出路径未覆盖（已跳过）：{'; '.join(mysql_skipped)}")
    if not mysql_passed and not mysql_skipped:
        print("MySQL 导出路径未覆盖：本次运行未包含 MySQL 用例")
    # 兜底：确认每个用例都被真正执行或显式跳过（真实失败会抛异常直接非零退出）。
    executed = set(passed) | {item.split("：", 1)[0] for item in skipped}
    if executed != {name for name, _ in checks}:
        raise SystemExit(1)
